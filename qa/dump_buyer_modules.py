import openpyxl

wb = openpyxl.load_workbook('qa/Mobile_App_Combined_Test_Results.xlsx', data_only=True)
ws = wb['All Results']
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
out = []
for r in range(2, ws.max_row + 1):
    row = {headers[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
    if not row.get('Test Case ID'):
        continue
    if row['Workbook'] != 'Buyer - Modules':
        continue
    out.append(
        f"=== {row['Test Case ID']} [{row['Sheet']}] {row['Status']}\n"
        f"Scenario: {row['Test Scenario']}\n"
        f"Expected: {row['Expected Result']}\n"
        f"Actual: {row['Actual Result']}\n"
        f"Remarks: {row['Remarks']}\n"
    )

with open('qa/buyer_modules_workbook.txt', 'w', encoding='utf-8', newline='\n') as f:
    f.write('\n'.join(out))
print(len(out))
