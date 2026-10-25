import sys
import openpyxl

path = sys.argv[1] if len(sys.argv) > 1 else 'qa/Mobile_App_Combined_Test_Results.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
print(wb.sheetnames)
for ws in wb.worksheets:
    print(ws.title, ws.max_row, ws.max_column)
    print([ws.cell(1, c).value for c in range(1, min(ws.max_column, 12) + 1)])

query = sys.argv[2] if len(sys.argv) > 2 else 'BUY-ACCT'
for ws in wb.worksheets:
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        joined = ' '.join(str(v) for v in row if v is not None)
        if query in joined:
            print('\n===', ws.title, 'row', r)
            for h, v in zip(headers, row):
                if v is not None:
                    print(f'{h}: {v}')
