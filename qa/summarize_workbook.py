import openpyxl
from collections import Counter, defaultdict

wb = openpyxl.load_workbook('qa/Mobile_App_Combined_Test_Results.xlsx', data_only=True)
ws = wb['All Results']
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
rows = []
for r in range(2, ws.max_row + 1):
    row = {headers[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
    if row.get('Test Case ID'):
        rows.append(row)

print('rows', len(rows))
print('status', Counter(r['Status'] for r in rows))
print('workbook', Counter(r['Workbook'] for r in rows))
print('sheet')
for k, v in sorted(Counter(r['Sheet'] for r in rows).items()):
    print(' ', k, v)

print('\nID prefix/status:')
by_prefix = defaultdict(Counter)
for r in rows:
    cid = str(r['Test Case ID'])
    prefix = '-'.join(cid.split('-')[:2])
    by_prefix[prefix][r['Status']] += 1
for prefix, c in sorted(by_prefix.items()):
    print(' ', prefix, dict(c))

print('\nNon-pass rows:')
for r in rows:
    if r['Status'] != 'Pass':
        print(f"{r['Test Case ID']}\t{r['Status']}\t{r['Sheet']}\t{r['Module']}\t{r['Test Scenario']}")
